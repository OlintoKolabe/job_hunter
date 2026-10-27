"""Mocked contract tests: no real API calls or billing."""
import io
import json
import os
import unittest
from datetime import date, timedelta
from unittest.mock import patch
from docx import Document
from ai_resume_matcher import evaluate, tailored_plan, original_evidence_plan, to_ranking, AIUnavailable
from resume_engine import audit, render_docx, EXPERIENCE

JD=("Senior Product Manager required to lead discovery, product roadmaps, stakeholders, "
    "backlog prioritization and Scrum delivery in a regulated business. "
    "Must have proven product ownership and leadership experience. "
    "Preferred: programming in Python and extensive Power BI dashboards. "
    "Own stakeholder alignment across teams, facilitate product discovery, translate "
    "customer research into roadmaps and collaborate with designers and engineers "
    "on measurable outcomes and continuous product delivery.")

class AIResumeMatcherTests(unittest.TestCase):
    def test_manual_outside_filters_overrides_eligibility_only(self):
        import app
        from fastapi import HTTPException

        base={"url":"https://example.com/jobs/123", "title":"Product Manager",
              "description":"A valid description "*25, "confirmed":False,
              "needs_verification":["pais"]}
        with patch.object(app,"require_auth",return_value=4), \
             patch.object(app.mu,"enabled",True), \
             patch.object(app.mu,"experience_integrity_issues",return_value=[]), \
             patch.object(app.mu,"factual_experience",return_value=[{"verified":True}]), \
             patch.object(app,"_cv_source",side_effect=RuntimeError("jd_lookup_reached")):
            with self.assertRaises(HTTPException) as blocked:
                app.ai_resume_download(app.ResumeAIRequest(**base),request=None)
            self.assertEqual(blocked.exception.status_code,422)

            with self.assertRaisesRegex(RuntimeError,"jd_lookup_reached"):
                app.ai_resume_download(app.ResumeAIRequest(**base,manual_outside_filters=True),request=None)

    def setUp(self):
        # Selection-contract tests do not hit the separate translation API.
        self._localize_patch=patch("ai_resume_matcher.localize_plan",side_effect=lambda p,j:p)
        self._localize_patch.start()
        self.addCleanup(self._localize_patch.stop)

    def test_jd_prose_determines_resume_language(self):
        from ai_resume_matcher import job_description_language
        english=("We are hiring a product manager with experience in product "
                 "development and stakeholder collaboration. You will work with "
                 "engineering teams and own the roadmap and delivery.")
        portuguese=("Buscamos profissional com experiência em gestão de produtos, "
                    "desenvolvimento de soluções e trabalho em equipe. "
                    "A empresa exige conhecimento de requisitos e responsabilidades.")
        self.assertEqual(job_description_language(english),"en")
        self.assertEqual(job_description_language(portuguese),"pt")

    def test_english_render_preserves_source_evidence_audit(self):
        from ai_resume_matcher import localize_plan
        import multiuser as mu
        profile={"name":"Candidate Test","email":"candidate@example.com",
                 "phone":"11987654321","location":"Brasil","linkedin":"",
                 "summary":"Atuei com produtos e equipes.","education":["Graduação — USP"],
                 "certifications":[],"languages":["Português (nativo)"],
                 "experience":[{"role":"Gerente de Produto","company":"Empresa A",
                     "dates":"jan 2022 - dez 2025",
                     "facts":["Gerenciei roadmap e backlog de produtos digitais."],
                     "verified":True}]}
        jd=("We are looking for an experienced Product Manager who will work "
            "with cross-functional teams and lead product roadmaps, collaboration "
            "and user research to deliver products to customers.")
        answer={"summary":"I worked with products and teams.",
                "entries":[{"role":"Product Manager","dates":"Jan 2022 - Dec 2025",
                            "facts":["Managed roadmaps and backlogs for digital products."]}],
                "education":["Undergraduate degree — USP"],
                "languages":["Portuguese (native)"]}
        class FakeResponse:
            def __enter__(self):return self
            def __exit__(self,*args):return False
            def read(self,*args):
                return json.dumps({"choices":[{"message":{"content":json.dumps(answer)}}],
                                   "usage":{"prompt_tokens":100,"completion_tokens":90}}).encode()
        token=mu.current_profile.set(profile)
        try:
            with patch.object(mu,"enabled",True),patch.dict(os.environ,{"OPENAI_API_KEY":"test","JH_AI_ENABLED":"1"}), \
                 patch("ai_resume_matcher._ask",return_value={"selected_fact_ids":["0:0"],"summary":"Product manager with documented experience."}), \
                 patch("ai_resume_matcher.urllib.request.urlopen",return_value=FakeResponse()):
                plan=tailored_plan(jd,"Product Manager")
                self._localize_patch.stop()
                import ai_resume_matcher as matcher
                matcher.localize_plan(plan,jd)
                document=render_docx(plan)
                self.assertTrue(audit(plan,document)["factual_bullets_verified"])
                text=" | ".join(p.text for p in Document(io.BytesIO(document)).paragraphs)
                self.assertIn("Professional Experience",text)
                self.assertIn("Managed roadmaps and backlogs",text)
                self.assertNotIn("Formação e certificações",text)
        finally:
            mu.current_profile.reset(token)

    def test_duplicate_facts_from_two_verified_roles_do_not_block_cv(self):
        import multiuser as mu
        facts="Led product discovery and roadmap prioritization with stakeholders."
        profile={"name":"Candidate Test","email":"candidate@example.com",
                 "phone":"11987654321","location":"Brasil","linkedin":"",
                 "summary":"","education":[],"certifications":[],"languages":[],
                 "experience":[
                   {"role":"Product Manager","company":"Company A","dates":"jan 2025 – atual",
                    "facts":[facts,"Delivered user research and prototyping for digital products."],"verified":True},
                   {"role":"Product Owner","company":"Company B","dates":"jan 2023 – dez 2024",
                    "facts":[facts,"Managed delivery with engineering and UX teams."],"verified":True}]}
        answer={"selected_fact_ids":["0:0","1:0","1:1","0:1"],"summary":"","rewrites":[]}
        token=mu.current_profile.set(profile)
        try:
            with patch.object(mu,"enabled",True),patch("ai_resume_matcher._ask",return_value=answer):
                plan=tailored_plan(JD,"Product Manager")
                doc=render_docx(plan)
                audit(plan,doc)
            chosen=[f for row in plan["experience"] for f in row["facts"]]
            self.assertEqual(chosen.count(facts),1)
            self.assertEqual(len(chosen),3)
        finally:
            mu.current_profile.reset(token)

    def test_real_jd_phrasing_is_not_lost_to_strict_word_overlap(self):
        # Source wording can include filler words even when a meaningful
        # quoted requirement is fully present in the job description.
        jd=('You will be responsible for the strategic management of our product '
            'roadmap and partnerships with key cross functional stakeholders. '
            'Prioritize the backlog with the engineering team and delivery managers.')
        answer={'score':72,'reasoning':'Evidence-based match',
                'requirements':[{'name':'strategic management of our product roadmap and partnerships',
                                 'mandatory':True,'match':'transferable','evidence_ids':['1:0']}]}
        with patch('ai_resume_matcher._ask',return_value=answer):
            result=evaluate(jd,'Product Manager')
        self.assertEqual(len(result['requirements']),1)
        self.assertEqual(result['requirements'][0]['match'],'transferable')

    def test_requirement_not_present_in_jd_is_never_credited(self):
        answer={'score':95,'requirements':[{'name':'Kubernetes platform administration',
                                          'mandatory':True,'match':'direct','evidence_ids':['1:0']}]}
        with patch('ai_resume_matcher._ask',return_value=answer):
            with self.assertRaisesRegex(ValueError,'requisitos auditáveis'):
                evaluate(JD,'Product Manager')

    def test_numeric_json_scores_are_normalized_without_inflation(self):
        base={'reasoning':'Relevant experience',
              'requirements':[{'name':'product roadmaps','mandatory':True,
                               'match':'direct','evidence_ids':['1:0']}]}
        for value,expected in [(75.4,75),('75',75),('75,4%',75),('75/100',75)]:
            with self.subTest(value=value),patch('ai_resume_matcher._ask',return_value={**base,'score':value}):
                self.assertEqual(evaluate(JD,'Product Manager')['score'],expected)
        for invalid in (None,True,-1,101,'very good','999/100'):
            with self.subTest(invalid=invalid),patch('ai_resume_matcher._ask',return_value={**base,'score':invalid}):
                with self.assertRaises(ValueError):evaluate(JD,'Product Manager')

    def test_no_key_no_paid_call(self):
        with patch.dict(os.environ,{"JH_AI_ENABLED":"0","OPENAI_API_KEY":""}):
            with self.assertRaises(AIUnavailable):
                evaluate(JD,"Senior Product Manager")

    def test_matcher_rejects_unproven_evidence_and_score_inflation(self):
        answer={"score":100,"reasoning":"Relevant PM experience, but gaps remain.",
                "requirements":[{"name":"product roadmaps","mandatory":True,"match":"direct","evidence_ids":["1:0"]},
                 {"name":"Python programming","mandatory":True,"match":"direct","evidence_ids":["999:9"]},
                 {"name":"Power BI dashboards","mandatory":False,"match":"unproven","evidence_ids":[]}]}
        with patch("ai_resume_matcher._ask",return_value=answer):
            result=evaluate(JD,"Senior Product Manager")
        self.assertEqual(result["requirements"][1]["match"],"unproven")
        self.assertLess(result["score"],80)
        ranking=to_ranking({"title":"Senior Product Manager","description":JD},"Product Manager",result)
        self.assertEqual(ranking["unverified_count"],2)
        self.assertEqual(ranking["mandatory_count"],2)

    def test_cv_is_ai_selected_and_does_not_invent_bullets(self):
        answer={"selected_fact_ids":["1:0","1:1","1:2","4:1","4:2"],
                "summary":"Profissional de produto com experiência em gestão de roadmaps e backlog."}
        with patch("ai_resume_matcher._ask",return_value=answer):
            plan=tailored_plan(JD,"Senior Product Manager","Empresa de testes")
        self.assertEqual(plan["cv_mode"],"ai_resume_matcher_evidence_v1")
        doc=render_docx(plan)
        self.assertTrue(audit(plan,doc)["factual_bullets_verified"])
        text=" ".join(x.text for x in Document(io.BytesIO(doc)).paragraphs)
        self.assertNotIn("Power BI dashboards",text)
        self.assertIn("Graduação",text)

    def test_tailored_cv_limits_secondary_roles_and_finishes_long_summary(self):
        long_summary="Delivered documented project work across teams. " * 19
        answer={"selected_fact_ids":["1:0","1:1","1:2","4:1"],
                "summary":long_summary}
        with patch("ai_resume_matcher._ask",return_value=answer):
            plan=tailored_plan(JD,"Technical Project Manager")
        self.assertLessEqual(len(plan["experience"])+len(plan["additional_experience"]),6)
        self.assertLessEqual(len(plan["additional_experience"]),2)
        self.assertNotIn("Professor de MBA",[x["role"] for x in plan["additional_experience"]])
        self.assertTrue(plan["summary"].endswith("."))
        self.assertLessEqual(len(plan["summary"]),520)
        self.assertTrue(audit(plan,render_docx(plan))["factual_bullets_verified"])

    def test_cv_rejects_sparse_or_unknown_selection(self):
        with patch("ai_resume_matcher._ask",return_value={"selected_fact_ids":["999:9","1:0"]}):
            plan=tailored_plan(JD,"Senior Product Manager")
        self.assertEqual(len(plan["experience"]),1)
        self.assertTrue(audit(plan,render_docx(plan))["factual_bullets_verified"])
        with patch("ai_resume_matcher._ask",return_value={"selected_fact_ids":["999:9"]}):
            with self.assertRaisesRegex(ValueError,"nenhuma evidência"):
                tailored_plan(JD,"Senior Product Manager")

    def test_rewritten_collision_recovers_only_original_verified_evidence(self):
        original=EXPERIENCE[1]["facts"][0]
        answer={"selected_fact_ids":["1:0","1:1"],
                "summary":"Gestão de produtos e resultados.",
                "rewrites":[{"id":"1:0","text":"Gestão e criação de roadmaps de produto, priorizando impacto e valor percebido."}]}
        with patch("ai_resume_matcher._ask",return_value=answer):
            plan=tailored_plan(JD,"Product Manager")
        recovered=original_evidence_plan(plan)
        self.assertIn(original,recovered["experience"][0]["facts"])
        self.assertEqual(recovered["bullet_provenance"],{})
        self.assertTrue(audit(recovered,render_docx(recovered))["factual_bullets_verified"])

    def test_ai_summary_repeats_bullet_and_verified_summary_restores_audit(self):
        original=EXPERIENCE[1]["facts"][0]
        with patch("ai_resume_matcher._ask",return_value={
                "selected_fact_ids":["1:0","1:1"],"summary":original}):
            plan=tailored_plan(JD,"Product Manager")
        with self.assertRaises(AssertionError):
            audit(plan,render_docx(plan))
        recovered=original_evidence_plan(plan)
        self.assertTrue(audit(recovered,render_docx(recovered))["factual_bullets_verified"])

    def test_ai_bullet_rewrites_require_original_fact_provenance(self):
        base={"selected_fact_ids":["1:0","1:1","1:2","4:1","4:2"],
              "summary":"Gestão de produtos, projetos e stakeholders.",
              "rewrites":[{"id":"1:0","text":"Gestão e criação de roadmaps de produto, priorizando impacto e valor percebido."},
                          {"id":"1:1","text":"Implementou SQL e Power BI para gerar ROI de 900% no departamento."}]}
        with patch("ai_resume_matcher._ask",return_value=base):
            plan=tailored_plan(JD,"Product Manager")
        facts=[f for entry in plan["experience"] for f in entry["facts"]]
        self.assertIn("Gestão e criação de roadmaps de produto, priorizando impacto e valor percebido.",facts)
        self.assertIn(EXPERIENCE[1]["facts"][1],facts)
        self.assertTrue(audit(plan,render_docx(plan))["factual_bullets_verified"])

    def test_post_filter_hydration_does_not_use_paid_matchmaking(self):
        from datetime import datetime,timezone
        from unittest.mock import Mock
        from jd_enrichment import hydrate_jobs
        url="https://example.com/job/1"
        cursor=Mock()
        cursor.fetchall.return_value=[(url,JD,datetime.now(timezone.utc))]
        ai={"score":68,"requirements":[{"name":"product roadmaps","mandatory":True,
             "match":"direct","evidence_ids":["1:0"]}],
             "reasoning":"Documented roadmaps","usage":{"prompt_tokens":100,"completion_tokens":80}}
        with patch.dict(os.environ,{"JH_AI_ENABLED":"1","OPENAI_API_KEY":"test-only","DATABASE_URL":"fake"}), \
             patch("jd_enrichment.get_cached",return_value=None), \
             patch("jd_enrichment.evaluate_ai",return_value=ai) as called, \
             patch("jd_enrichment.reserve",return_value=9), \
             patch("jd_enrichment.settle") as settled, \
             patch("jd_enrichment.save") as saved, \
             patch("jd_enrichment.psycopg.connect") as conn:
            result,_=hydrate_jobs(cursor,[{"url":url,"title":"Senior Product Manager",
                                           "description":JD}],"Product Manager",lambda x:x)
        self.assertIsNone(result[0]["ranking"]["score"])
        called.assert_not_called()
        settled.assert_not_called()
        saved.assert_not_called()
        self.assertNotEqual(result[0]["ai_status"],"interpreted")

    def test_authenticated_route_returns_actual_docx_not_chatgpt_url(self):
        import asyncio
        from unittest.mock import Mock
        import app as service
        result={"selected_fact_ids":["1:0","1:1","1:2","4:1","4:2"],
                "summary":"Experiência comprovada em produtos, roadmaps e gestão de equipes."}
        async def drain(response):
            return b"".join([part async for part in response.body_iterator])
        with patch.dict(os.environ,{"JH_AI_ENABLED":"1","OPENAI_API_KEY":"test-only"}), \
             patch("app.require_auth"), patch("app._cv_source",return_value=("example.com/job/1",JD)), \
             patch("score_repository.reserve",return_value=23), \
             patch("score_repository.settle"),patch("app.psycopg.connect"), \
             patch("ai_resume_matcher._ask",return_value=result):
            payload=service.ResumeAIRequest(url="https://example.com/job/1",title="Senior Product Manager",company="Example")
            response=service.ai_resume_download(payload,Mock())
            document=asyncio.run(drain(response))
        self.assertEqual(response.status_code,200)
        self.assertEqual(document[:2],b"PK")
        self.assertIn("wordprocessingml",response.media_type)
        self.assertIn("AI=ai_resume_matcher_evidence_v1",
                      Document(io.BytesIO(document)).core_properties.keywords)

    def test_download_recovers_when_ai_summary_duplicates_evidence(self):
        from unittest.mock import Mock
        import app as service
        selected_fact=EXPERIENCE[1]["facts"][0]
        answer={"selected_fact_ids":["1:0","1:1"],"summary":selected_fact}
        with patch.dict(os.environ,{"JH_AI_ENABLED":"1","OPENAI_API_KEY":"test-only"}), \
             patch("app.require_auth"),patch("app._cv_source",return_value=("example.com/job/1",JD)), \
             patch("score_repository.reserve",return_value=41),patch("score_repository.settle"), \
             patch("app.psycopg.connect"),patch("ai_resume_matcher._ask",return_value=answer):
            response=service.ai_resume_download(service.ResumeAIRequest(
                url="https://example.com/job/1",title="Product Manager"),Mock())
            document=__import__("asyncio").run(self._drain(response))
        content=" ".join(p.text for p in Document(io.BytesIO(document)).paragraphs)
        self.assertEqual(content.count(selected_fact),1)
        self.assertEqual(response.status_code,200)

    def test_cv_uses_full_listing_description_when_public_page_is_blocked(self):
        import asyncio
        from unittest.mock import Mock
        import app as service
        answer={"selected_fact_ids":["1:0","1:1","1:2","4:1","4:2"],"summary":"Gestão de produtos e projetos."}
        with patch.dict(os.environ,{"JH_AI_ENABLED":"1","OPENAI_API_KEY":"test-only"}), \
             patch("app.require_auth"),patch("app._cv_source",side_effect=service.HTTPException(404,"no cache")), \
             patch("score_repository.reserve",return_value=35),patch("score_repository.settle"), \
             patch("app.psycopg.connect"),patch("ai_resume_matcher._ask",return_value=answer), \
             patch("jd_retrieval.extract_public_jd") as public:
            response=service.ai_resume_download(service.ResumeAIRequest(
                url="https://br.indeed.com/viewjob?jk=245b508b66b7f110",
                title="Senior Product Manager",description=JD*2),Mock())
            payload=asyncio.run(self._drain(response))
        public.assert_not_called()
        self.assertEqual(payload[:2],b"PK")

    @staticmethod
    async def _drain(response):
        return b"".join([chunk async for chunk in response.body_iterator])

    def test_failed_cv_without_provider_response_releases_full_reservation(self):
        from unittest.mock import Mock
        import app as service
        with patch.dict(os.environ,{"JH_AI_ENABLED":"1","OPENAI_API_KEY":"test-only"}), \
             patch("app.require_auth"), patch("app._cv_source",return_value=("example.com/job/1",JD)), \
             patch("score_repository.reserve",return_value=24), \
             patch("score_repository.settle") as settled, \
             patch("app.psycopg.connect"), \
             patch("ai_resume_matcher._ask",side_effect=ValueError("invalid response")):
            with self.assertRaises(service.HTTPException) as raised:
                service.ai_resume_download(service.ResumeAIRequest(
                    url="https://example.com/job/1",title="Senior Product Manager"),Mock())
        self.assertEqual(raised.exception.status_code,422)
        settled.assert_called_once_with(unittest.mock.ANY,24,0.0,False)

    def test_failed_cv_after_provider_response_charges_only_actual_tokens(self):
        from unittest.mock import Mock
        import app as service
        from ai_resume_matcher import _last_billable_usage, estimate_cost
        usage={"prompt_tokens":4000,"completion_tokens":800}
        def incomplete_response(*_args):
            _last_billable_usage.set(usage)
            return {"selected_fact_ids":["999:9"]}
        with patch.dict(os.environ,{"JH_AI_ENABLED":"1","OPENAI_API_KEY":"test-only"}), \
             patch("app.require_auth"), patch("app._cv_source",return_value=("example.com/job/1",JD)), \
             patch("score_repository.reserve",return_value=25), \
             patch("score_repository.settle") as settled, \
             patch("app.psycopg.connect"), \
             patch("ai_resume_matcher._ask",side_effect=incomplete_response):
            with self.assertRaises(service.HTTPException):
                service.ai_resume_download(service.ResumeAIRequest(
                    url="https://example.com/job/1",title="Senior Product Manager"),Mock())
        settled.assert_called_once_with(unittest.mock.ANY,25,estimate_cost(usage),False)

class ShortlistTests(unittest.TestCase):
    def test_http_shortlist_rejects_oversized_batch(self):
        from unittest.mock import Mock
        from app import ranking_shortlist, ShortlistRequest
        from fastapi import HTTPException
        with patch('app.require_auth'):
            with self.assertRaises(HTTPException) as error:
                ranking_shortlist(Mock(),ShortlistRequest(jobs=[{}]*2501,roles='Product Manager'))
        self.assertEqual(error.exception.status_code,400)

    def test_endpoint_never_calls_paid_ai(self):
        from unittest.mock import Mock
        from app import ranking_shortlist, ShortlistRequest
        with patch('app.require_auth'),patch('ai_resume_matcher._ask',side_effect=AssertionError('Unexpected paid AI')):
            result=ranking_shortlist(Mock(),ShortlistRequest(
                jobs=[{'url':'https://example.org/v', 'title':'Product Manager'}],
                roles='Product Manager',limit=10))
        self.assertEqual(result['urls'],['https://example.org/v'])
        self.assertEqual(result['ai_calls'],0)

    def test_global_relevance_and_diversity_without_tokens(self):
        from ai_resume_matcher import shortlist_jobs
        rows = [dict(url=f'https://example.org/jobs/{i}', title='Office Assistant') for i in range(150)]
        rows.append(dict(url='https://example.org/jobs/best', title='Senior Product Manager',
                         description='roadmap product strategy backlog agile stakeholder ' * 25))
        rows.append(dict(url='https://example.org/jobs/project', title='Project Manager'))
        chosen = shortlist_jobs(rows, 'Product Manager, Project Manager', 10)
        self.assertIn('https://example.org/jobs/best', chosen[:2])
        self.assertIn('https://example.org/jobs/project', chosen[:2])
        self.assertEqual(len(chosen), 10)

    def test_global_shortlist_excludes_unverified_and_duplicates(self):
        from ai_resume_matcher import shortlist_jobs
        rows = [dict(url='https://a.com/a', title='Product Manager', needs_verification=['pais']),
                dict(url='https://b.com/b', title='Product Manager', source='indeed'),
                dict(url='https://b.com/b', title='Product Manager', source='ashby')]
        self.assertEqual(shortlist_jobs(rows, 'Product Manager'), ['https://b.com/b'])

class ShortlistedJDTests(unittest.TestCase):
    def test_inhire_full_jd_uses_existing_provider_api_for_cv_and_ranking(self):
        from jd_enrichment import fallback_public_jd
        job={'url':'https://cobli.inhire.app/vagas/cf249d92-5a0c-4143-bb3a-c7e5ea688d00/product-manager-s-nior-account-experience',
             'title':'Product Manager Sênior','company':'Cobli','source':'inHire'}
        body='Responsabilidades: liderar roadmap e stakeholders. '*12
        with patch('jd_enrichment.fetch_json',return_value={'jobPage':{'description':body}}) as api, \
             patch('jd_enrichment.extract_public_jd',side_effect=AssertionError('No fallback needed')):
            self.assertEqual(fallback_public_jd(job),body)
        self.assertEqual(api.call_args.args[0],
                         'https://api.inhire.app/job-posts/public/pages/cf249d92-5a0c-4143-bb3a-c7e5ea688d00')
        self.assertEqual(api.call_args.args[1],{'X-Tenant':'cobli'})

    def test_disabled_ai_matchmaking_does_not_charge_budget(self):
        from jd_enrichment import hydrate_jobs
        from ai_resume_matcher import _last_billable_usage, estimate_cost
        from unittest.mock import Mock
        usage={'prompt_tokens':4600,'completion_tokens':700}
        def invalid_ai(*args):
            _last_billable_usage.set(usage)
            raise ValueError('IA não apresentou requisitos auditáveis')
        cursor=Mock()
        cursor.fetchall.return_value=[]
        job={'url':'https://example.org/job','title':'Product Manager',
             'description':JD}
        with patch.dict(os.environ,{'JH_AI_ENABLED':'1','OPENAI_API_KEY':'test','DATABASE_URL':'test'}), \
             patch('jd_enrichment.get_cached',return_value=None), \
             patch('jd_enrichment.evaluate_ai',side_effect=invalid_ai), \
             patch('jd_enrichment.reserve',return_value=9), \
             patch('jd_enrichment.settle') as settle, \
             patch('jd_enrichment.psycopg.connect'):
            result,_=hydrate_jobs(cursor,[job],'Product Manager',lambda url:url)
        self.assertIsNone(result[0]['ranking']['score'])
        settle.assert_not_called()

    def test_indeed_ranking_reuses_existing_exact_jobspy_jd(self):
        from jd_enrichment import fallback_public_jd
        job={'url':'https://br.indeed.com/viewjob?jk=ed26935acd1ff08e',
             'title':'Product Manager','company':'Empresa'}
        full='Responsible for product discovery and roadmaps. '*12
        with patch('indeed_provider.description_for_job',return_value=full) as spy, \
             patch('jd_enrichment.extract_public_jd',side_effect=AssertionError('unnecessary scraper')):
            self.assertEqual(fallback_public_jd(job),full)
        spy.assert_called_once_with(job['url'],job['title'],job['company'])

    def test_failed_indeed_jobspy_can_fall_back_to_public_html(self):
        from jd_enrichment import fallback_public_jd
        job={'url':'https://br.indeed.com/viewjob?jk=e41955656030e3b2',
             'title':'Product Manager','company':'Empresa'}
        with patch('indeed_provider.description_for_job',return_value=None), \
             patch('jd_enrichment.extract_public_jd',return_value='public JD') as html_reader:
            self.assertEqual(fallback_public_jd(job),'public JD')
        html_reader.assert_called_once_with(job['url'])

    def test_other_sources_do_not_run_indeed_search(self):
        from jd_enrichment import fallback_public_jd
        job={'url':'https://boards.greenhouse.io/example/jobs/123','title':'Product Manager'}
        with patch('indeed_provider.description_for_job',side_effect=AssertionError('Indeed only')), \
             patch('jd_enrichment.extract_public_jd',return_value='generic JD'):
            self.assertEqual(fallback_public_jd(job),'generic JD')

    def test_filter_does_not_destroy_collector_jd_before_cache(self):
        from datetime import datetime, timezone
        from jd_enrichment import persist_listing_jds
        from search_filters import analyze_filters
        from unittest.mock import Mock
        jd='Product strategy, roadmap, discovery, requirements and stakeholder leadership. '*12
        job={'url':'https://example.com/product/1','title':'Senior Product Manager',
             'company':'Example','location':'Remote Brazil','remote':True,
             'source':'ashby','published':datetime.now(timezone.utc).isoformat(),'description':jd}
        confirmed,review,_=analyze_filters([job],'Product Manager','BR','remote',3,'newest')
        self.assertEqual(len(confirmed),1)
        self.assertEqual(confirmed[0]['description'],'')
        cur=Mock()
        stats=persist_listing_jds(cur,[job],confirmed+review,lambda u:u)
        self.assertEqual(stats['ashby']['jd_valid'],1)
        self.assertEqual(cur.executemany.call_args.args[1][0][1],jd)

    def test_missing_or_challenge_jd_never_enters_cache(self):
        from jd_enrichment import valid_listing_jd,persist_listing_jds
        from unittest.mock import Mock
        self.assertFalse(valid_listing_jd('Verify you are human. '*90))
        job={'url':'https://example.com/product/2','source':'inhire','description':''}
        cur=Mock()
        stats=persist_listing_jds(cur,[job],[job],lambda u:u)
        self.assertEqual(stats['inhire']['jd_missing'],1)
        cur.executemany.assert_not_called()

    def test_prefilter_cache_is_off_by_default_and_never_queries(self):
        from jd_enrichment import cached_jds_for_filter
        original=[{'url':'https://example.com/1','title':'Product Manager','description':''}]
        with patch.dict(os.environ,{'JH_PREFILTER_JD_CACHE_ENABLED':'0'}), \
             patch('jd_enrichment.psycopg.connect',side_effect=AssertionError('cache must stay off')):
            self.assertIs(cached_jds_for_filter(original,lambda url:url,'unused'),original)

    def test_prefilter_cache_preserves_collector_fields_and_valid_jd(self):
        from jd_enrichment import cached_jds_for_filter
        from unittest.mock import MagicMock
        full='Remote Brazil. The position requires product roadmap ownership and agile delivery. '*10
        collector='Product leadership, agile delivery and planning with stakeholders. '*10
        original=[{'url':'https://example.com/1','title':'Product Manager','company':'A',
                   'location':'Remote Brazil','remote':True,'description':''},
                  {'url':'https://example.com/2','title':'Project Manager','company':'B',
                   'location':'Remote Brazil','description':collector}]
        connection=MagicMock()
        connection.__enter__.return_value.cursor.return_value.__enter__.return_value.fetchall.return_value=[
            ('https://example.com/1',full)]
        with patch.dict(os.environ,{'JH_PREFILTER_JD_CACHE_ENABLED':'1'}), \
             patch('jd_enrichment.psycopg.connect',return_value=connection):
            result=cached_jds_for_filter(original,lambda url:url,'unused')
        self.assertEqual(result[0]['description'],full)
        self.assertEqual(result[0]['title'],original[0]['title'])
        self.assertEqual(result[0]['location'],original[0]['location'])
        self.assertTrue(result[0]['remote'])
        self.assertEqual(result[1],original[1])
        self.assertEqual(original[0]['description'],'')
        self.assertEqual(connection.__enter__.return_value.cursor.return_value.__enter__.return_value.execute.call_count,1)

    def test_prefilter_cache_failure_leaves_filter_input_unchanged(self):
        from jd_enrichment import cached_jds_for_filter
        import psycopg
        original=[{'url':'https://example.com/1','title':'Product Manager','description':''}]
        with patch.dict(os.environ,{'JH_PREFILTER_JD_CACHE_ENABLED':'1'}), \
             patch('jd_enrichment.psycopg.connect',side_effect=psycopg.OperationalError('offline')):
            self.assertIs(cached_jds_for_filter(original,lambda url:url,'unused'),original)

if __name__=="__main__":unittest.main()
