import io
import unittest
from docx import Document
from resume_engine import build_plan,render_docx,audit,EXPERIENCE,CONTACT

JD="We need a business analyst who manages stakeholder needs, roadmaps and backlog in Azure DevOps, works with product owners and defines user stories. Previous hands-on Oracle programming experience is a plus."

class ResumeEngineTests(unittest.TestCase):
    def test_multiuser_docx_never_mixes_sidebar_into_education(self):
        import multiuser as mu
        from unittest.mock import patch
        profile={"name":"Marcus Pablo Rezende Teixeira Manso Pereira",
                 "location":"Santos, SP","phone":"13 99212-8060",
                 "email":"marcus@example.com","linkedin":"linkedin.com/in/marcus-pablo",
                 "summary":"I turn raw data into strategic decisions — and manual processes\\\ninto intelligent automation. &amp; analytics.",
                 "experience":[{"role":"Automation Analyst","company":"Bluetech Films",
                                "dates":"June 2025 - August 2025",
                                "facts":["Built Python and SQL pipelines for analytics."],"verified":True}],
                 "education":["Pontifícia Universidade Católica de Goiás", "Engenharia de Computação",
                              "Marcus Pablo Rezende Teixeira Manso Pereira", "Python",
                              "Data Scientist & Data Engineer"],
                 "certifications":["Asimov Academy","Senior Associate @PwC"],
                 "skills":["Python"]}
        token=mu.current_profile.set(profile)
        try:
            with patch.object(mu,"enabled",True):
                plan=build_plan("Python SQL ETL analytics", "Data Engineer")
                text=[p.text for p in Document(io.BytesIO(render_docx(plan))).paragraphs]
                self.assertEqual(text[0],profile["name"])
                self.assertEqual(sum(profile["name"]==line for line in text),1)
                self.assertNotIn("Data Scientist & Data Engineer", text)
                self.assertNotIn("Senior Associate @PwC",text)
                self.assertIn("Engenharia de Computação", text)
                self.assertTrue(any("intelligent automation. & analytics." in line for line in text))
        finally:
            mu.current_profile.reset(token)

    def test_docx_has_plain_ats_typography_without_blue_headings_or_title_rule(self):
        from docx.shared import RGBColor
        doc=Document(io.BytesIO(render_docx(build_plan(JD,"Senior Business Analyst"))))
        for style_name in ("Title","Heading 1","Heading 2"):
            self.assertEqual(doc.styles[style_name].font.color.rgb,RGBColor(0,0,0))
        self.assertLessEqual(doc.styles["Title"].font.size.pt,20)
        self.assertGreaterEqual(doc.styles["Normal"].font.size.pt,10)
        self.assertEqual(len(doc.sections),1)
        self.assertEqual(len(doc.tables),0)
        self.assertTrue(audit(build_plan(JD,"Senior Business Analyst"),
                              render_docx(build_plan(JD,"Senior Business Analyst")))["factual_bullets_verified"])

    def test_document_provenance_survives_download_and_fallback(self):
        from resume_engine import CV_DOCUMENT_VERSION
        for mode in ("hybrid","cached_hybrid","fallback","cached_fallback","deterministic"):
            plan=build_plan(JD,"Senior Business Analyst")
            plan["cv_mode"]=mode
            plan["cv_build"]="abcdef123456"
            doc=Document(io.BytesIO(render_docx(plan)))
            tag=doc.core_properties.keywords
            self.assertIn("JobHunterCV="+CV_DOCUMENT_VERSION,tag)
            self.assertIn("AI="+mode,tag)
            self.assertIn("Build=abcdef123456",tag)
            self.assertIn("JD="+plan["jd_sha256"][:16],tag)
            self.assertTrue(audit(plan,render_docx(plan))["factual_bullets_verified"])
    def test_dentistry_never_appears_even_when_jd_requests_it(self):
        from resume_engine import EDUCATION
        self.assertFalse(any("odontolog" in degree.casefold() for degree in EDUCATION))
        for title,jd in [
            ("Dentista","Formação em Odontologia e atendimento odontológico"),
            ("Product Manager","Roadmap product discovery stakeholders"),
            ("Analyst","CRM growth operations")
        ]:
            plan=build_plan(jd,title)
            self.assertFalse(any("odontolog" in degree.casefold() for degree in plan["education"]))
            doc=Document(io.BytesIO(render_docx(plan)))
            body=" ".join(p.text for p in doc.paragraphs)
            self.assertNotIn("Graduação em Odontologia",body)
            self.assertTrue(audit(plan,render_docx(plan))["factual_bullets_verified"])

    def test_tailored_cv_omits_irrelevant_dentistry_degree_without_inventing_degree(self):
        for role in ("Product Manager","Business Analyst","Marketing Manager"):
            plan=build_plan("Roadmap CRM business analysis and stakeholder planning",role)
            text=" ".join(p.text for p in Document(io.BytesIO(render_docx(plan))).paragraphs)
            self.assertNotIn("Graduação em Odontologia",text)
            self.assertIn("MBA em Marketing",text)

    def test_ai_workflow_requirement_uses_documented_tools_not_invented_expertise(self):
        jd="Senior Product Manager: product discovery, roadmap, project delivery, stakeholders; proactively apply AI-driven tools in product workflows."
        p=build_plan(jd,"Senior Product Manager")
        chosen=[f for e in p["experience"] for f in e["facts"]]
        self.assertTrue(any("Codex" in f and "Claude Code" in f for f in chosen))
        rendered=" ".join(x.text for x in Document(io.BytesIO(render_docx(p))).paragraphs)
        self.assertNotIn("engenheiro de inteligência artificial",rendered.lower())

    def test_business_analysis_does_not_promote_generic_marketing_role(self):
        p=build_plan("Analyze business requirements, stakeholder needs, backlog and user stories; develop functional specifications.","Business Analyst")
        self.assertNotIn("Aviva",{e["company"] for e in p["experience"]})

    def test_encoded_job_html_does_not_rank_benefit_boilerplate(self):
        from resume_engine import relevant_jd_text
        html_jd="&lt;div&gt;Company introduction.&lt;/div&gt;&lt;p&gt;Role expectations: Deliver roadmaps, product discovery and stakeholder alignment.&lt;/p&gt;&lt;p&gt;Why work for us? Unlimited fruit and snacks.&lt;/p&gt;"
        extracted=relevant_jd_text(html_jd)
        self.assertIn("product discovery",extracted)
        self.assertNotIn("fruit and snacks",extracted)
        self.assertNotIn("Company introduction",extracted)

    def test_reopened_docx_technical_grading(self):
        from resume_engine import evaluate_rendered_resume
        plan=build_plan(JD,"Senior Business Analyst")
        data=render_docx(plan)
        result=evaluate_rendered_resume(plan,data,JD)
        self.assertTrue(result["factual_ok"])
        self.assertIn("factual_traceability",result["dimensions"])
        self.assertEqual(result["jd_sha256"],plan["jd_sha256"])
        self.assertGreater(result["technical_score"],0)

    def test_complete_linkedin_history_adds_marketing_evidence(self):
        from resume_engine import EXPERIENCE
        self.assertTrue(any("R$ 7 milhões" in fact for fact in EXPERIENCE[12]["facts"]))
        self.assertTrue(any("triplicada" in fact for fact in EXPERIENCE[9]["facts"]))
        self.assertTrue(any("Hotjar" in fact for fact in EXPERIENCE[1]["facts"]))

    def test_program_manager_gets_governance_evidence(self):
        from resume_engine import EXPERIENCE
        jd="Program Manager governance portfolio cross-functional teams risk project delivery budget"
        p=build_plan(jd,"Program Manager")
        detail={e["source_index"]:e["facts"] for e in p["experience"]}
        self.assertIn(2,detail)
        self.assertIn(8,detail)
        self.assertTrue(any("Governança corporativa" in f for f in detail[8]))
        self.assertTrue(audit(p,render_docx(p))["factual_bullets_verified"])
    def test_performance_media_docx_surfaces_marketing_career(self):
        jd=("Lead digital performance campaigns, paid media, media budgets, "
            "Google Analytics, CRM, growth, audience targeting and A/B testing.")
        plan=build_plan(jd,"Coordenador de Mídia De Performance Digital")
        detailed={e["source_index"] for e in plan["experience"]}
        self.assertIn(9,detailed)
        self.assertIn("liderança de marketing",plan["summary"])
        data=render_docx(plan)
        self.assertTrue(audit(plan,data)["factual_bullets_verified"])
        paragraphs=[p.text for p in Document(io.BytesIO(data)).paragraphs]
        self.assertTrue(any("Gestão de marketing, campanhas omnichannel, CRM" in p for p in paragraphs))

    def test_marketing_manager_highlights_earlier_marketing_leadership(self):
        p=build_plan("Marketing strategy CRM campaigns branding and marketing budget","Marketing Manager")
        detailed={e["source_index"] for e in p["experience"]}
        self.assertIn(9,detailed)
        self.assertTrue({9,11,12,13}&detailed)
    def test_generates_document_without_ai(self):
        p=build_plan(JD,"Senior Business Analyst")
        b=render_docx(p)
        a=audit(p,b)
        self.assertEqual(a["cost_usd"],0)
        self.assertEqual(a["ai_calls"],0)
        self.assertTrue(a["factual_bullets_verified"])
    def test_contact_and_chronology(self):
        p=build_plan(JD,"Senior Business Analyst")
        self.assertEqual([x["company"] for x in p["experience"]+p["additional_experience"]].__len__(),len(EXPERIENCE))
        from resume_engine import chronology_index
        self.assertEqual([chronology_index(EXPERIENCE[e["source_index"]]) for e in p["experience"]], sorted([chronology_index(EXPERIENCE[e["source_index"]]) for e in p["experience"]],reverse=True))
        doc=Document(io.BytesIO(render_docx(p)))
        content="\n".join(x.text for x in doc.paragraphs)
        for c in CONTACT.values():self.assertIn(c,content)
    def test_no_unverified_technical_claim(self):
        p=build_plan(JD,"Senior Business Analyst")
        b=render_docx(p)
        content="\n".join(x.text for x in Document(io.BytesIO(b)).paragraphs)
        self.assertNotIn("Oracle programming",content)
        self.assertNotIn("PowerBuilder",content)
    def test_reproducible_bytes_and_hash(self):
        p=build_plan(JD,"Senior Business Analyst")
        a=audit(p,render_docx(p))
        b=audit(p,render_docx(p))
        self.assertEqual(a["jd_sha256"],b["jd_sha256"])
    def test_no_duplicate_bullets_or_highlight_section(self):
        p=build_plan(JD,"Senior Business Analyst")
        doc=Document(io.BytesIO(render_docx(p)))
        text="\n".join(x.text for x in doc.paragraphs)
        self.assertNotIn("Competências e experiências relevantes para a vaga",text)
        bullets=[f for e in p["experience"] for f in e["facts"]]
        self.assertEqual(len(bullets),len(set(bullets)))
        self.assertTrue(all(text.count(f)==1 for f in bullets))
    def test_distinct_jds_select_distinct_relevant_evidence(self):
        a=build_plan("Business analyst backlog requirements user stories stakeholders Azure DevOps","Analyst")
        b=build_plan("Marketing digital paid ads CRM churn vendas growth leads aquisição","Growth")
        fa=[f for e in a["experience"] for f in e["facts"]]
        fb=[f for e in b["experience"] for f in e["facts"]]
        self.assertNotEqual(fa,fb)
        self.assertNotEqual(a["summary"],b["summary"])
    def test_document_maintains_earlier_roles_compact(self):
        p=build_plan(JD,"Senior Business Analyst")
        self.assertTrue(p["additional_experience"])
        self.assertTrue(all("facts" not in e for e in p["additional_experience"]))
    def test_single_chronological_career_timeline(self):
        from resume_engine import chronology_index
        p=build_plan(JD,"Business Analyst")
        parsed=Document(io.BytesIO(render_docx(p)))
        text="\\n".join(x.text for x in parsed.paragraphs)
        self.assertNotIn("Outras experiências",text)
        order=sorted(EXPERIENCE,key=chronology_index,reverse=True)
        positions=[text.find(e["role"]+" — "+e["company"]) for e in order]
        self.assertTrue(all(position>=0 for position in positions))
        self.assertEqual(positions,sorted(positions))

    def test_internal_docx_homologation_matrix(self):
        """Run the actual engine and reopen DOCX for each distinct JD category."""
        from resume_engine import categories, TOPICS, FACT_TOPICS
        from docx import Document
        matrix = [
            ("Business Analyst", "Business Analyst: map stakeholder needs, document user stories, prioritize backlog, support Scrum and Azure DevOps.", "analysis"),
            ("Product Owner", "Product Owner: own roadmap, discovery, backlog, user research and alignment with stakeholders.", "product"),
            ("Growth Manager", "Growth Manager: CRM, retention, acquisition, marketing campaigns, churn, leads, Google Ads.", "growth"),
            ("Project Manager", "Project Manager: portfolio governance, delivery, resources, budget, and processes.", "project"),
            ("Marketing Manager", "Marketing Manager: coordinate marketing strategy, branding, paid campaigns, CRM and sales.", "growth"),
            ("Data Product Manager", "Data Product Manager: roadmap, KPIs, indicators, analytics and product discovery.", "data"),
        ]
        seen=set()
        for title,jd,required in matrix:
            with self.subTest(title=title):
                plan=build_plan(jd,title)
                content=render_docx(plan)
                report=audit(plan,content)
                doc=Document(io.BytesIO(content))
                body=" ".join(p.text for p in doc.paragraphs)
                facts=[f for e in plan["experience"] for f in e["facts"]]
                covered=set().union(*(categories(f,FACT_TOPICS) for f in facts))
                self.assertTrue(report["factual_bullets_verified"])
                self.assertIn(required,covered, f"No supported evidence of {required}: {title}")
                self.assertIn("Resumo profissional",body)
                self.assertIn("Experiência profissional",body)
                self.assertGreaterEqual(len(facts),3)
                self.assertNotIn(report["jd_sha256"],seen)
                seen.add(report["jd_sha256"])

    def test_product_output_is_scored_from_reopened_docx(self):
        """Inspect the actual DOCX paragraphs, not the selection plan alone."""
        import re
        from resume_engine import TOPICS, FACT_TOPICS, categories
        cases=[
            ("Business Analyst", "Translate product strategy and roadmap priorities into acceptance journeys and test cases; gather stakeholder requirements; optimize business processes; support UAT, release outcomes and acceptance criteria.",
             {"analysis","product","project"}),
            ("Product Owner", "Prioritize roadmap, product discovery, user research, backlog and stakeholder alignment; own outcomes and sprint planning.",
             {"product","agility"}),
            ("Marketing Manager", "Lead CRM campaigns, brand strategy, customer retention, acquisition, marketing budget and commercial teams.",
             {"growth","people"}),
            ("Project Manager", "Manage project portfolio, governance, resources, budgets, delivery and multi-disciplinary teams.",
             {"project","people"}),
        ]
        for title,jd,expected in cases:
            with self.subTest(title=title):
                plan=build_plan(jd,title)
                output=render_docx(plan)
                self.assertTrue(audit(plan,output)["factual_bullets_verified"])
                doc=Document(io.BytesIO(output))
                rendered=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
                content=" ".join(rendered)
                self.assertEqual(rendered.count("Resumo profissional"),1)
                self.assertEqual(rendered.count("Experiência profissional"),1)
                facts=[p.text for p in doc.paragraphs if p.style.name=="List Bullet"]
                cv_evidence=[x for x in facts if x in {f for e in EXPERIENCE for f in e["facts"]}]
                self.assertEqual(len(cv_evidence),len(set(cv_evidence)))
                self.assertGreaterEqual(len(cv_evidence),3)
                topics=set().union(*(categories(f,FACT_TOPICS) for f in cv_evidence))
                coverage=len(topics & expected)/len(expected)
                self.assertGreaterEqual(coverage,0.5,
                    f"Rendered CV coverage too low: {title}, {coverage:.0%}")
                # Explicitly disallow industry/technology claims found in the JD
                # if the underlying career evidence does not contain them.
                for unsupported in ("Oracle programming","PowerBuilder","life insurance claims","UAT"):
                    if unsupported.casefold() not in " ".join(f for e in EXPERIENCE for f in e["facts"]).casefold():
                        self.assertNotIn(unsupported.casefold(),content.casefold())

    def test_education_contains_usp_and_mbas_without_degree_area(self):
        from resume_engine import EDUCATION
        text=" | ".join(EDUCATION)
        self.assertIn("Graduação — Universidade de São Paulo (USP)",text)
        self.assertIn("MBA em Marketing",text)
        self.assertIn("MBA em Administração e Gestão",text)
        self.assertIn("PSPO I",text)
        self.assertNotIn("Odontologia",text)

    def test_empty_jd_rejected(self):
        with self.assertRaises(ValueError):build_plan("  ","A")

class BackendResumeRouteContractTests(unittest.TestCase):
    """Only the new GPT/MCP path can produce a user-selected ATS DOCX."""
    def test_verified_gpt_download_reopens_and_audits_docx(self):
        import app, asyncio, hashlib
        from unittest.mock import patch
        from itsdangerous import URLSafeTimedSerializer
        from resume_engine import EXPERIENCE
        jd="Product Manager roadmap discovery stakeholders backlog delivery"
        evidence=["0:0"]
        signer=URLSafeTimedSerializer("test-only-secret")
        payload={"url":"https://example.com/vaga","title":"Product Manager",
                 "company":"Empresa","ids":evidence,
                 "jd_hash":hashlib.sha256(jd.encode()).hexdigest()}
        ticket=signer.dumps(payload,salt="cv-action-download-v1")
        with patch.object(app,"serializer",signer),patch.object(app,"_cv_source",return_value=(payload["url"],jd)):
            response=app.gpt_cv_download(ticket)
        async def gather():return b"".join([part async for part in response.body_iterator])
        parsed=Document(io.BytesIO(asyncio.run(gather())))
        content=" ".join(p.text for p in parsed.paragraphs)
        self.assertIn(EXPERIENCE[0]["facts"][0],content)
        self.assertIn("gpt_action_verified_facts",parsed.core_properties.keywords)
        self.assertEqual(response.headers["x-cv-mode"],"gpt_action_verified_facts")
        self.assertTrue(response.headers["content-disposition"].endswith(".docx"))
    def test_unverified_gpt_claim_is_rejected(self):
        import app,hashlib
        from unittest.mock import patch
        from itsdangerous import URLSafeTimedSerializer
        from fastapi import HTTPException
        jd="Project Manager"
        signer=URLSafeTimedSerializer("test-only-secret")
        payload={"url":"https://example.com/vaga","title":"Project Manager",
                 "company":"Empresa","ids":["999:999"],
                 "jd_hash":hashlib.sha256(jd.encode()).hexdigest()}
        ticket=signer.dumps(payload,salt="cv-action-download-v1")
        with patch.object(app,"serializer",signer),patch.object(app,"_cv_source",return_value=(payload["url"],jd)):
            with self.assertRaises(HTTPException) as error:app.gpt_cv_download(ticket)
        self.assertEqual(error.exception.status_code,400)
    def test_legacy_resume_generate_route_removed(self):
        import app
        paths=[route.path for route in app.app.routes if hasattr(route,"path")]
        self.assertNotIn("/api/resume/generate",paths)

if __name__=="__main__":unittest.main()


class GPTActionContractTests(unittest.TestCase):
    def test_only_private_mcp_and_signed_download_exist(self):
        import app
        paths=[route.path for route in app.app.routes if hasattr(route,"path")]
        self.assertIn("/api/gpt-action/download",paths)
        self.assertNotIn("/api/gpt-action/context",paths)
        self.assertNotIn("/api/gpt-action/cv",paths)
        self.assertNotIn("/api/resume/gpt-start",paths)
        self.assertFalse(hasattr(app,"_cv_action_auth"))

    def test_action_export_does_not_use_llm_key(self):
        import inspect,app
        source=inspect.getsource(app.gpt_cv_download)
        self.assertIn("audit(plan,data)",source)
        self.assertIn("max_age=900",source)
        self.assertNotIn("apply_hybrid(",source)


    def test_mcp_uses_one_docx_renderer_and_bearer_header(self):
        import inspect,app
        source=inspect.getsource(app.prepare_verified_cv)
        self.assertNotIn("render_docx(",source)
        self.assertNotIn("access_token",inspect.signature(app.fetch_cv_job).parameters)
        self.assertNotIn("access_token",inspect.signature(app.prepare_verified_cv).parameters)
        from unittest.mock import patch
        import asyncio
        class Inner:
            async def __call__(self,scope,receive,send):
                raise AssertionError("Unauthenticated MCP request reached inner server")
        messages=[]
        async def receive():return {"type":"http.request","body":b"","more_body":False}
        async def send(message):messages.append(message)
        with patch.dict("os.environ",{"CV_ACTION_TOKEN":"test-secret"}):
            asyncio.run(app._CVAuthorizedMCP(Inner())({"type":"http","headers":[],"method":"POST","path":"/"},
                  receive,send))
        self.assertEqual(messages[0]["status"],401)

    def test_gpt_source_is_verified_only(self):
        from resume_engine import EXPERIENCE
        ids={f"{i}:{j}" for i,e in enumerate(EXPERIENCE) for j,_ in enumerate(e["facts"])}
        self.assertTrue(ids)
        self.assertNotIn("999:999",ids)
